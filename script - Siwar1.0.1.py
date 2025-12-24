import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
from PIL import Image, ImageEnhance, ImageFilter, ImageTk
import pytesseract
import numpy as np
from docx import Document
from pdf2image import convert_from_path
import threading
import ttkbootstrap as ttk
from ttkbootstrap.constants import *

# ====== Paths & Tesseract setup ======
if getattr(sys, 'frozen', False):
    base_path = sys._MEIPASS
else:
    base_path = os.path.dirname(os.path.abspath(__file__))

tess_path = os.path.join(base_path, "Tesseract-OCR")
pytesseract.pytesseract.tesseract_cmd = os.path.join(tess_path, "tesseract.exe")
os.environ['TESSDATA_PREFIX'] = os.path.join(tess_path, "tessdata")

custom_config = r'--oem 3 --psm 6'
language = "fra+ara"  # French + Arabic

# ====== OCR Functions ======
def deskew_image(np_image):
    binary = np_image > 0
    coords = np.column_stack(np.where(binary))
    if coords.size == 0:
        return Image.fromarray(np_image)
    cov = np.cov(coords.T)
    evals, evecs = np.linalg.eigh(cov)
    principal_vector = evecs[:, np.argsort(evals)[::-1][0]]
    angle = np.arctan2(principal_vector[1], principal_vector[0]) * (180.0 / np.pi)
    pil_image = Image.fromarray(np_image)
    return pil_image.rotate(-angle, expand=True, fillcolor=255)

def preprocess_image(image_path):
    image = Image.open(image_path)
    return preprocess_image_from_pil(image)

def preprocess_image_from_pil(image):
    gray = image.convert("L")
    gray = ImageEnhance.Contrast(gray).enhance(2)
    gray = gray.point(lambda x: 0 if x < 140 else 255, '1')
    gray = gray.filter(ImageFilter.SHARPEN)
    np_image = np.array(gray).astype(np.uint8) * 255
    gray = deskew_image(np_image)
    return gray

def correct_orientation(image):
    try:
        osd = pytesseract.image_to_osd(image)
        rotate = 0
        for line in osd.split("\n"):
            if "Rotate" in line:
                rotate = int(line.split(":")[1].strip())
        if rotate != 0:
            image = image.rotate(-rotate, expand=True, fillcolor=255)
    except:
        pass
    return image

def pdf_to_images(pdf_path, dpi=300):
    if getattr(sys, 'frozen', False):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    poppler_path = os.path.join(base_path, "Poppler", "bin")
    return convert_from_path(pdf_path, dpi=dpi, poppler_path=poppler_path)

def ocr_images(folder_path, update_progress_callback=None):
    texts = []
    files = sorted([f for f in os.listdir(folder_path) if f.lower().endswith((".jpg", ".jpeg", ".png", ".pdf"))])
    total = len(files)

    for i, filename in enumerate(files, start=1):
        file_path = os.path.join(folder_path, filename)

        if filename.lower().endswith(".pdf"):
            images = pdf_to_images(file_path)
            first_page_image = images[0]
            text = ""
            for page_num, image in enumerate(images, start=1):
                preprocessed_image = preprocess_image_from_pil(image)
                preprocessed_image = correct_orientation(preprocessed_image)
                page_text = pytesseract.image_to_string(preprocessed_image, lang=language, config=custom_config)
                text += f"\n--- Page {page_num} ---\n{page_text}"
            texts.append((filename, file_path, first_page_image, text))
        else:
            preprocessed_image = preprocess_image(file_path)
            preprocessed_image = correct_orientation(preprocessed_image)
            text = pytesseract.image_to_string(preprocessed_image, lang=language, config=custom_config)
            texts.append((filename, file_path, preprocessed_image, text))

        if update_progress_callback:
            update_progress_callback(i, total)

    return texts

# ====== Modern OCR GUI ======
class OCRApp:
    def __init__(self, root):
        self.root = root
        self.root.geometry("950x700")
        self.root.title("Siwar Image Reader v1.0.1")
        self.root.minsize(900, 650)

        style = ttk.Style("flatly")

        # Icon
        icon_path = os.path.join(base_path, "icon.ico")
        if os.path.exists(icon_path):
            self.root.iconbitmap(icon_path)

        self.folder_path = ""
        self.texts = []
        self.current_index = 0
        self.loading_index = 0
        self.loading_animation_running = False

        # Menu
        menu = tk.Menu(root)
        root.config(menu=menu)
        help_menu = tk.Menu(menu, tearoff=0)
        help_menu.add_command(label="About", command=self.show_about)
        menu.add_cascade(label="Help", menu=help_menu)

        # Frames
        self.top_frame = ttk.Frame(root, padding=10)
        self.top_frame.pack(fill="x")
        self.mid_frame = ttk.Frame(root, padding=10)
        self.mid_frame.pack(fill="both", expand=True)
        self.bottom_frame = ttk.Frame(root, padding=10)
        self.bottom_frame.pack(fill="x")

        # Top Frame: Folder + Read
        self.folder_btn = ttk.Button(self.top_frame, text="Select Folder", bootstyle=INFO, command=self.select_folder)
        self.folder_btn.pack(side=LEFT, padx=5)
        self.read_btn = ttk.Button(self.top_frame, text="Read Files", bootstyle=SUCCESS, command=self.read_files)
        self.read_btn.pack(side=LEFT, padx=5)
        self.progress_label = ttk.Label(self.top_frame, text="", bootstyle=PRIMARY)
        self.progress_label.pack(side=LEFT, padx=20)

        # Middle Frame: Image Preview + Text
        self.preview_frame = ttk.Frame(self.mid_frame)
        self.preview_frame.pack(side=LEFT, fill="both", expand=True)

        self.original_label = ttk.Label(self.preview_frame, text="Original", anchor="center", bootstyle=SECONDARY)
        self.original_label.pack(fill="both", expand=True, padx=5, pady=5)
        self.processed_label = ttk.Label(self.preview_frame, text="Preprocessed / PDF Preview", anchor="center", bootstyle=SECONDARY)
        self.processed_label.pack(fill="both", expand=True, padx=5, pady=5)

        self.text_frame = ttk.Frame(self.mid_frame)
        self.text_frame.pack(side=LEFT, fill="both", expand=True, padx=5)
        self.text_area = ScrolledText(self.text_frame, font=("Segoe UI", 11))
        self.text_area.pack(fill="both", expand=True)

        # Bottom Frame: Navigation
        self.prev_btn = ttk.Button(self.bottom_frame, text="<< Prev", bootstyle=WARNING, command=self.prev_text, state=DISABLED)
        self.prev_btn.pack(side=LEFT, padx=5)
        self.next_btn = ttk.Button(self.bottom_frame, text="Next >>", bootstyle=WARNING, command=self.next_text, state=DISABLED)
        self.next_btn.pack(side=LEFT, padx=5)

    # ====== Animated loading ======
    def start_loading_animation(self):
        self.loading_index = 0
        self.loading_animation_running = True
        self.animate_loading()

    def animate_loading(self):
        if not self.loading_animation_running:
            return
        dots = '.' * (self.loading_index % 4)
        self.text_area.delete("1.0", tk.END)
        self.text_area.insert(tk.END, f"Processing files{dots}\n")
        self.loading_index += 1
        self.root.after(500, self.animate_loading)

    def stop_loading_animation(self):
        self.loading_animation_running = False

    # ====== Other methods ======
    def show_about(self):
        messagebox.showinfo(
            "About",
            "Siwar Image Reader v1.0.1\n\n"
            "Developer: Walid Ouadfeul\n"
            "Contact: walid.ouadfeul@gmail.com\n\n"
            "Supports French and Arabic OCR for images and PDFs,\n"
            "and extracts the text to TXT and Word files."
        )

    def select_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.folder_path = folder
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, f"Selected folder:\n{folder}\n")

    def update_progress(self, current, total):
        self.progress_label.config(text=f"Processing file {current} of {total}")
        self.root.update()

    def read_files(self):
        if not self.folder_path:
            messagebox.showwarning("Warning", "Please select a folder first!")
            return

        self.read_btn.config(state=DISABLED)
        self.prev_btn.config(state=DISABLED)
        self.next_btn.config(state=DISABLED)

        self.start_loading_animation()

        # Run OCR in a background thread
        def ocr_task():
            self.texts = ocr_images(self.folder_path, update_progress_callback=self.update_progress)
            self.stop_loading_animation()
            self.root.after(100, self.show_results)

        threading.Thread(target=ocr_task, daemon=True).start()

    def show_results(self):
        if self.texts:
            # Save TXT
            txt_file = os.path.join(self.folder_path, "extracted_textconv.txt")
            with open(txt_file, "w", encoding="utf-8") as f:
                for filename, _, _, text in self.texts:
                    f.write(f"--- Text from {filename} ---\n{text}\n\n")

            # Save DOCX
            docx_file = os.path.join(self.folder_path, "extracted_textconv.docx")
            doc = Document()
            with open(txt_file, "r", encoding="utf-8") as file:
                for line in file:
                    doc.add_paragraph(line.strip())
            doc.save(docx_file)

            messagebox.showinfo("Done", f"OCR complete!\nText saved to:\n{txt_file}\n{docx_file}")
            self.current_index = 0
            self.show_text()
        else:
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, "No images or PDFs found in folder.")

        self.read_btn.config(state=NORMAL)

    def show_text(self):
        filename, img_path, preview_image, text = self.texts[self.current_index]

        if preview_image:
            display_img = preview_image.copy()
            display_img.thumbnail((400, 400))
            self.tk_processed = ImageTk.PhotoImage(display_img)
            self.processed_label.config(image=self.tk_processed, text="Preprocessed / PDF Preview", compound="top")

            if img_path.lower().endswith((".jpg", ".jpeg", ".png")):
                original_img = Image.open(img_path)
                original_img.thumbnail((400, 400))
                self.tk_original = ImageTk.PhotoImage(original_img)
                self.original_label.config(image=self.tk_original, text="Original Image", compound="top")
            else:
                original_img = preview_image.copy()
                original_img.thumbnail((400, 400))
                self.tk_original = ImageTk.PhotoImage(original_img)
                self.original_label.config(image=self.tk_original, text="PDF First Page", compound="top")
        else:
            self.original_label.config(image="", text="N/A", compound="top")
            self.processed_label.config(image="", text="N/A", compound="top")

        self.text_area.delete(1.0, tk.END)
        self.text_area.insert(tk.END, f"--- Text from {filename} ---\n{text}")

        self.prev_btn.config(state=NORMAL if self.current_index > 0 else DISABLED)
        self.next_btn.config(state=NORMAL if self.current_index < len(self.texts) - 1 else DISABLED)
        self.progress_label.config(text=f"File {self.current_index + 1} of {len(self.texts)}")

    def prev_text(self):
        if self.texts and self.current_index > 0:
            self.current_index -= 1
            self.show_text()

    def next_text(self):
        if self.texts and self.current_index < len(self.texts) - 1:
            self.current_index += 1
            self.show_text()


# ====== Run App ======
if __name__ == "__main__":
    root = ttk.Window(title="Siwar Image Reader v1.0.1", themename="flatly")
    app = OCRApp(root)
    root.mainloop()
